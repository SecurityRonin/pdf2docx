"""A ruled table must become a real Word table, not flowing paragraphs.

Flattening tables was the largest remaining cause of page drift: a table
occupying a third of a page became text and the document grew by pages.
"""
import re
import zipfile

import docx

from pdf2docx import convert
from pdf2docx.tables import page_tables


def _w(path):
    return zipfile.ZipFile(path).read("word/document.xml").decode("utf8")


def test_a_ruled_table_is_detected(tabled_pdf):
    tables = page_tables(tabled_pdf)
    assert tables, "no table detected"
    t = tables[0]
    assert t.nrows >= 3 and t.ncols >= 2
    assert any(c.text for c in t.cells)


def test_table_becomes_a_real_word_table(tabled_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(tabled_pdf, out)
    xml = _w(out)
    assert "<w:tbl>" in xml, "table must be a real w:tbl"
    assert xml.count("<w:tr") >= 3
    assert "gridCol" in xml


def test_table_cells_hold_their_text(tabled_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(tabled_pdf, out)
    d = docx.Document(str(out))
    assert d.tables, "python-docx sees no table"
    cell_text = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Alice" in cell_text and "Analyst" in cell_text


def test_table_text_is_not_duplicated_as_paragraphs(tabled_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(tabled_pdf, out)
    body = " ".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Analyst" not in body, "table text leaked into flowing paragraphs"


def test_header_row_is_repeated_across_pages(tabled_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(tabled_pdf, out)
    assert "tblHeader" in _w(out)


def test_merged_cell_becomes_grid_span(tabled_pdf, tmp_path):
    """A cell spanning two columns must survive as gridSpan, not two cells."""
    out = tmp_path / "out.docx"
    convert(tabled_pdf, out)
    assert "gridSpan" in _w(out)
