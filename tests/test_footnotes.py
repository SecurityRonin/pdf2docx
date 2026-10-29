"""A footnote must be a real Word footnote: a reference run in the body, and its
text in footnotes.xml — not the note appended as prose at the end.
"""
import pathlib
import zipfile

import docx

from pdf2docx import convert

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def test_footnote_text_moves_into_footnotes_part(footnote_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(footnote_pdf, out)
    with zipfile.ZipFile(out) as z:
        assert "word/footnotes.xml" in z.namelist()
        notes = z.read("word/footnotes.xml").decode("utf8")
    assert "Statement of A. Example" in notes


def test_body_carries_a_footnote_reference(footnote_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(footnote_pdf, out)
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode("utf8")
    assert "footnoteReference" in body


def test_note_text_is_not_left_in_the_body(footnote_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(footnote_pdf, out)
    d = docx.Document(str(out))
    body_text = " ".join(p.text for p in d.paragraphs)
    assert "Statement of A. Example" not in body_text


def test_the_note_marker_is_not_left_in_body_text(footnote_pdf, tmp_path):
    """The superscript '1' must become a reference, not literal characters."""
    out = tmp_path / "out.docx"
    convert(footnote_pdf, out)
    d = docx.Document(str(out))
    body_text = " ".join(p.text for p in d.paragraphs)
    assert body_text.rstrip().endswith("12 March.")
