"""Bullets and numbered items must be real Word list paragraphs backed by
numbering.xml — not lines that merely look like a list.
"""
import pathlib
import zipfile

import docx

from pdf2docx import convert

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _list_paras(path):
    d = docx.Document(str(path))
    return [p for p in d.paragraphs if p.style.name.startswith("List")]


def test_bulleted_lines_become_list_paragraphs(numbered_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(numbered_pdf, out)
    texts = [p.text.strip() for p in _list_paras(out)]
    assert "Mail server" in texts
    assert "File server" in texts
    assert "Directory service" in texts


def test_list_paragraphs_reference_numbering(numbered_pdf, tmp_path):
    """A 'List' style with no numPr prints without a bullet in Word."""
    out = tmp_path / "out.docx"
    convert(numbered_pdf, out)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf8")
    assert "numPr" in xml, "list paragraphs must carry native numbering"
    assert "numId" in xml


def test_the_bullet_glyph_is_not_kept_in_the_text(numbered_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(numbered_pdf, out)
    for p in _list_paras(out):
        assert not p.text.strip().startswith(("•", "-", "*")), p.text


def test_list_items_are_not_styled_as_headings(numbered_pdf, tmp_path):
    out = tmp_path / "out.docx"
    convert(numbered_pdf, out)
    d = docx.Document(str(out))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Mail server" not in heads
